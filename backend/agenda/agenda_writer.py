"""Write engine results into the manager's Meeting Agenda workbook.

Opens the agenda template with openpyxl, writes the FY25 figures the 2025 CSV
can prove (monthly + Jan->month accumulated sales, profit margins, profit
amounts) into their exact cells, and preserves the sheet's existing variance
formulas (=D6-C6, =E6/C6, =I6-H6, =J6/H6). When no 2026 CSV is supplied the
FY26 columns are left blank and their headers are annotated "awaiting data
input" (SKILL.md section 5 -- never fabricate 2026 numbers).

openpyxl does not evaluate formulas, so the preserved variance cells keep their
old cached results. We set fullCalcOnLoad so Excel/LibreOffice recompute them
when the file is opened -- the "recalc pass" referenced in SKILL.md section 7.
"""
from __future__ import annotations
import calendar

import openpyxl
from openpyxl.styles import Font

from .engine import TARGET_BRANCHES, branch_figures

# Variance % colouring: negative → red, zero/positive → blue (bold, matching template).
_RED  = Font(bold=True, color="FF0000")
_BLUE = Font(bold=True, color="0070C0")


def _pct_font(value: float | None):
    """Return red Font for negative values, blue Font for zero/positive."""
    if value is None:
        return None
    return _RED if value < 0 else _BLUE

# The six target branches sit two rows apart: an even "sales" row and the
# "Profit" row beneath it. JCI=6/7, KMT=8/9, GPL=10/11, MRT=12/13, MFW=14/15,
# SAT=16/17 -- see data/reference/Meeting_Agenda.xlsx.
SALES_ROW = {b: 6 + 2 * i for i, b in enumerate(TARGET_BRANCHES)}

AWAITING = "awaiting data input"

# FY26 cells cleared when no 2026 data exists. Keyed by (column, row offset
# from the branch sales row): 0 = sales row, 1 = profit row beneath it.
FY26_CELLS = [
    ("D", 0),  # FY26 monthly sales
    ("D", 1),  # FY26 monthly margin
    ("G", 1),  # FY26 monthly profit amount
    ("E", 1),  # FY26 monthly margin variance (value, not a formula)
    ("I", 0),  # FY26 accumulated sales
    ("I", 1),  # FY26 accumulated margin
    ("L", 1),  # FY26 accumulated profit amount
    ("J", 1),  # FY26 accumulated margin variance (value, not a formula)
]


def _acc_margin_label(branch: str, margin: float) -> str:
    """Reproduce the sheet's 'JCI        37.40%' accumulated-margin text."""
    return f"{branch}        {margin * 100:.2f}%"


def _write_branch(ws, branch: str, df25, df26, month: int,
                  df25_acc=None, acc25: dict | None = None) -> None:
    srow = SALES_ROW[branch]
    prow = srow + 1
    months = list(range(1, month + 1))

    monthly = branch_figures(df25, branch, [month])

    # --- FY25 monthly block (left side) ---
    ws[f"C{srow}"] = monthly["sales"]
    ws[f"C{prow}"] = monthly["margin"]
    ws[f"G{srow}"] = monthly["profit"]

    # --- FY25 accumulated (XLS dict takes priority over CSV DataFrame) ---
    if acc25 is not None and branch in acc25:
        d = acc25[branch]
        acc_sales, acc_profit, acc_margin = d["sales"], d["profit"], d["margin"]
    elif df25_acc is not None:
        acc = branch_figures(df25_acc, branch, months)
        acc_sales, acc_profit, acc_margin = acc["sales"], acc["profit"], acc["margin"]
    else:
        acc_sales = acc_profit = acc_margin = None

    if acc_sales is not None:
        ws[f"H{srow}"] = acc_sales
        ws[f"H{prow}"] = _acc_margin_label(branch, acc_margin)
        ws[f"L{srow}"] = acc_profit
    else:
        ws[f"H{srow}"] = ws[f"H{prow}"] = ws[f"L{srow}"] = None

    # --- FY26 columns ---
    if df26 is not None:
        m26 = branch_figures(df26, branch, [month])
        a26 = branch_figures(df26, branch, months)
        ws[f"D{srow}"] = m26["sales"]
        ws[f"D{prow}"] = m26["margin"]
        ws[f"G{prow}"] = m26["profit"]

        # Monthly margin variance (profit row, col E) — value cell, colour by sign.
        e_val = round(m26["margin"] - monthly["margin"], 4)
        ws[f"E{prow}"] = e_val
        ws[f"E{prow}"].font = _pct_font(e_val)

        ws[f"I{srow}"] = a26["sales"]
        ws[f"I{prow}"] = a26["margin"]
        ws[f"L{prow}"] = a26["profit"]

        # Accumulated margin variance (profit row, col J) — value cell, colour by sign.
        j_val = round(a26["margin"] - acc_margin, 4) if acc_margin is not None else None
        ws[f"J{prow}"] = j_val
        ws[f"J{prow}"].font = _pct_font(j_val)

        # Monthly sales variance % (sales row, col F) — formula cell, colour by sign.
        if monthly["sales"]:
            ws[f"F{srow}"].font = _pct_font(
                (m26["sales"] - monthly["sales"]) / monthly["sales"]
            )

        # Accumulated sales variance % (sales row, col K) — formula cell, colour by sign.
        if acc_sales and acc_sales != 0:
            ws[f"K{srow}"].font = _pct_font(
                (a26["sales"] - acc_sales) / acc_sales
            )
    else:
        for col, off in FY26_CELLS:
            ws[f"{col}{srow + off}"] = None


_UNSET = object()


def fill_agenda(template_path: str, out_path: str, df25, df26=None,
                month: int = 4, df25_acc=_UNSET,
                acc25: dict | None = None) -> str:
    """Fill a copy of the agenda template and save it to out_path.

    df25/df26: DataFrames from load_report_csv. df26 may be None.
    df25_acc:  accumulated FY25 DataFrame source (defaults to df25). Pass None
               when only a monthly CSV is available and no XLS is provided.
    acc25:     pre-decoded XLS branch totals from load_xls_report(); takes
               priority over df25_acc for the FY25 accumulated cells.
    The =D6-C6 style variance formulas are never touched.
    """
    if df25_acc is _UNSET:
        df25_acc = df25

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Sheet1"]

    month_name = calendar.month_name[month].upper()
    month_abbr = calendar.month_abbr[month].upper()
    ws["F1"] = month_name
    has_acc25 = acc25 is not None or df25_acc is not None
    ws["H2"] = f"JAN-{month_abbr} 25" if has_acc25 else f"JAN-{month_abbr} 25 ({AWAITING})"
    ws["I2"] = f"JAN-{month_abbr} 26"
    ws["H4"] = f"G 1-{month}"
    ws["I4"] = f"H 1-{month}"

    for branch in TARGET_BRANCHES:
        _write_branch(ws, branch, df25, df26, month, df25_acc, acc25)

    if df26 is None:
        ws["D5"] = f"FY26/ ({AWAITING})"
        ws["I5"] = f"FY26/ Acc. Sales ({AWAITING})"

    # Recalc pass: openpyxl cannot evaluate formulas in-process, so flag the
    # workbook for a full recalculation when a spreadsheet app opens it.
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return out_path
